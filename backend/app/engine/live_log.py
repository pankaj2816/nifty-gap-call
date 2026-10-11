"""Live gap-test log from 12 Oct 2026.

The prediction is locked from the completed US session only. Indian open and
close are written later as the result. A locked call is never rewritten.
"""
from datetime import date, datetime, time, timedelta, timezone
import html
import json
import threading
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.core.market_calendar import CalendarUnavailable, nse_session_bounds, nse_session_to_describe
from app.core.time_utils import IST, MARKET_CLOSE_TIME, MARKET_OPEN_TIME, now_ist
from app.engine.overnight import describe_overnight

LIVE_START = date(2026, 10, 12)
LOG_PATH = Path(__file__).resolve().parents[2] / ".." / "live_log" / "nifty_gap_log.xlsx"
_LOCK = threading.Lock()

COLUMNS = [
    ("session_date", "Session date"),
    ("locked_at_ist", "Prediction locked (IST)"),
    ("us_session", "US session"),
    ("spx_move_pct", "S&P move %"),
    ("gap_prediction", "Gap prediction"),
    ("status", "Status"),
    ("prev_close", "Previous NIFTY close"),
    ("nifty_open", "NIFTY open"),
    ("gap_pct", "Gap %"),
    ("gap_outcome", "Gap outcome"),
    ("nifty_close", "NIFTY close"),
    ("open_to_close_pct", "Open to close %"),
    ("open_to_close_outcome", "Open-to-close outcome"),
    ("india_vix", "India VIX known before the open"),
    ("prior_nifty_day_pct", "Previous NIFTY day %"),
    ("review", "What the result shows"),
    ("notes", "Notes"),
]
NOTE = "The call uses the completed US session only. Indian prices are filled afterwards and are not used to choose the call."


def _sessions_through(last: date) -> List[date]:
    days = []
    day = LIVE_START
    while day <= last:
        try:
            bounds = nse_session_bounds(day)
        except CalendarUnavailable:
            break
        if bounds:
            days.append(day)
        day += timedelta(days=1)
    return days


def _direction(later: float, earlier: float) -> str:
    if later > earlier:
        return "UP"
    if later < earlier:
        return "DOWN"
    return "FLAT"


def _close_is_final(bar: Dict[str, Any], now: datetime, fetched_at: Optional[float]) -> bool:
    if bar.get("is_complete") is False:
        return False
    if bar.get("is_complete") is True:
        return True
    day = date.fromisoformat(bar["date"])
    if day < now.astimezone(IST).date():
        return True
    if now < datetime.combine(day, MARKET_CLOSE_TIME, IST):
        return False
    if fetched_at is None:
        return False
    fetched = datetime.fromtimestamp(fetched_at, timezone.utc).astimezone(IST)
    return fetched >= datetime.combine(day, MARKET_CLOSE_TIME, IST)


def _blank_row(session: date) -> Dict[str, Any]:
    return {
        "session_date": session.isoformat(),
        "locked_at_ist": "",
        "us_session": "",
        "spx_move_pct": "",
        "gap_prediction": "",
        "status": "Waiting for the US close",
        "prev_close": "",
        "nifty_open": "",
        "gap_pct": "",
        "gap_outcome": "Waiting",
        "nifty_close": "",
        "open_to_close_pct": "",
        "open_to_close_outcome": "Waiting",
        "india_vix": "",
        "prior_nifty_day_pct": "",
        "review": "",
        "notes": NOTE,
    }


def _lock_prediction(row: Dict[str, Any], session: date, spx_bars, spx_stale: bool, prev_close) -> None:
    if row.get("locked_at_ist"):
        return
    clock = datetime.combine(session, time(8, 0), IST)
    read = describe_overnight(spx_bars, clock, is_stale=spx_stale)
    if read.get("session_date") != session.isoformat() or read.get("strength") in ("pending", "stale", "unavailable"):
        row["status"] = "Waiting for the US close" if read.get("strength") == "pending" else "Waiting for US data"
        return
    call = read.get("gap_call") or "NO_CALL"
    row.update(
        locked_at_ist=clock.strftime("%Y-%m-%d %H:%M IST"),
        us_session=read.get("us_date") or "",
        spx_move_pct=read.get("us_return_pct", ""),
        gap_prediction="NO CALL" if call == "NO_CALL" else call,
        status="Prediction locked",
        prev_close=prev_close if prev_close is not None else "",
    )


def _score_row(row: Dict[str, Any], session: date, nifty_bars, now: datetime, fetched_at: Optional[float]) -> None:
    if not row.get("locked_at_ist"):
        return
    earlier = [b for b in nifty_bars if b.get("date", "") < session.isoformat()]
    _known_context(row, earlier)
    if earlier and row.get("prev_close") in ("", None):
        row["prev_close"] = earlier[-1]["close"]
    bar = next((b for b in nifty_bars if b.get("date") == session.isoformat()), None)
    opened = now >= datetime.combine(session, MARKET_OPEN_TIME, IST)
    if not bar or not opened or row.get("prev_close") in ("", None):
        if row["gap_prediction"] == "NO CALL":
            row["gap_outcome"] = "No call — not scored"
            row["open_to_close_outcome"] = "No call — not scored"
            _review(row)
        return
    prev = float(row["prev_close"])
    row["nifty_open"] = bar["open"]
    gap_pct = (bar["open"] / prev - 1.0) * 100.0
    row["gap_pct"] = round(gap_pct, 3)
    if row["gap_prediction"] == "NO CALL":
        row["gap_outcome"] = "No call — not scored"
        row["status"] = "No call"
    else:
        gap_dir = _direction(bar["open"], prev)
        row["gap_outcome"] = "Flat" if gap_dir == "FLAT" else ("Same way" if gap_dir == row["gap_prediction"] else "Opposite")
        row["status"] = "Gap scored"
    if not _close_is_final(bar, now, fetched_at):
        _review(row)
        return
    row["nifty_close"] = bar["close"]
    row["open_to_close_pct"] = round((bar["close"] / bar["open"] - 1.0) * 100.0, 3)
    if row["gap_prediction"] == "NO CALL":
        row["open_to_close_outcome"] = "No call — not scored"
        row["status"] = "No call"
        _review(row)
        return
    day_dir = _direction(bar["close"], bar["open"])
    row["open_to_close_outcome"] = "Flat" if day_dir == "FLAT" else ("Same way" if day_dir == row["gap_prediction"] else "Opposite")
    row["status"] = "Day scored"
    _review(row)


def _known_context(row: Dict[str, Any], earlier: List[Dict[str, Any]]) -> None:
    """Facts available before the Indian open. Later prices must not replace them."""
    if earlier and row.get("india_vix") in ("", None) and earlier[-1].get("vix") not in (None, ""):
        row["india_vix"] = earlier[-1]["vix"]
    if len(earlier) >= 2 and row.get("prior_nifty_day_pct") in ("", None):
        row["prior_nifty_day_pct"] = round((earlier[-1]["close"] / earlier[-2]["close"] - 1.0) * 100.0, 3)


def _review(row: Dict[str, Any]) -> None:
    """Describe the scored result. This does not change the locked call."""
    if not row.get("locked_at_ist"):
        return
    call = row.get("gap_prediction")
    if call == "NO CALL":
        extra = ""
        if row.get("gap_pct") not in ("", None):
            extra = f" The open later moved {row['gap_pct']}% versus the previous close. This day is not counted."
        row["review"] = (
            f"No call. The S&P move was {row.get('spx_move_pct')}%, under the 1% line.{extra}"
        )
        return
    if row.get("gap_outcome") in ("", "Waiting"):
        return
    parts = []
    if row["gap_outcome"] == "Opposite":
        parts.append(
            f"The opening gap went the other way ({row.get('gap_pct')}%). "
            f"The call used only the S&P move of {row.get('spx_move_pct')}% and did not see the Indian open."
        )
    elif row["gap_outcome"] == "Same way":
        parts.append(f"The opening gap matched ({row.get('gap_pct')}%).")
    elif row["gap_outcome"] == "Flat":
        parts.append("The open was unchanged from the previous close.")
    if row.get("open_to_close_outcome") == "Opposite":
        parts.append(
            f"After the open, the session went the other way ({row.get('open_to_close_pct')}%). "
            "The rule is about the gap, so this is recorded separately and does not rewrite the call."
        )
    elif row.get("open_to_close_outcome") == "Same way":
        parts.append(f"After the open, the session also went the same way ({row.get('open_to_close_pct')}%).")
    try:
        us_move = abs(float(row.get("spx_move_pct")))
    except (TypeError, ValueError):
        us_move = None
    if row["gap_outcome"] == "Opposite" and us_move is not None and us_move < 1.5:
        parts.append(
            "The US move was only just past 1%. Days like this stay in the sheet so a later review can see whether small moves miss more often. One day does not change the rule."
        )
    if row.get("india_vix") not in ("", None):
        parts.append(f"India VIX known before the open was {row['india_vix']}.")
    if row.get("prior_nifty_day_pct") not in ("", None):
        parts.append(f"NIFTY's previous session moved {row['prior_nifty_day_pct']}%.")
    row["review"] = " ".join(parts)


def _summary(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    calls = [r for r in rows if r.get("gap_prediction") in ("UP", "DOWN")]
    gap_scored = [r for r in calls if r.get("gap_outcome") in ("Same way", "Opposite", "Flat")]
    day_scored = [r for r in calls if r.get("open_to_close_outcome") in ("Same way", "Opposite", "Flat")]
    return {
        "sessions": len(rows),
        "calls": len(calls),
        "no_calls": sum(r.get("gap_prediction") == "NO CALL" for r in rows),
        "gap_scored": len(gap_scored),
        "gap_same_way": sum(r.get("gap_outcome") == "Same way" for r in gap_scored),
        "day_scored": len(day_scored),
        "day_same_way": sum(r.get("open_to_close_outcome") == "Same way" for r in day_scored),
        "waiting": sum(not r.get("locked_at_ist") for r in rows),
    }


def _read_rows(path: Path) -> Dict[str, Dict[str, Any]]:
    if not path.exists():
        return {}
    book = load_workbook(path)
    sheet = book["Gap test"] if "Gap test" in book.sheetnames else book.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    title_to_key = {title: key for key, title in COLUMNS}
    found = {}
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        if not cells or not cells[0]:
            continue
        row = _blank_row(date.today())
        for title, value in zip(header, cells):
            key = title_to_key.get(title)
            if key:
                row[key] = "" if value is None else value
        found[str(row["session_date"])] = row
    return found


def _write(path: Path, rows: List[Dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Gap test"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for col, (key, title) in enumerate(COLUMNS, 1):
        cell = sheet.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = 64 if key == "review" else 18 if col < len(COLUMNS) else 72
    for r, row in enumerate(rows, 2):
        for c, (key, _) in enumerate(COLUMNS, 1):
            sheet.cell(r, c, row.get(key, ""))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    about = book.create_sheet("How to read this")
    about["A1"] = "Live test from 12 October 2026"
    about["A1"].font = Font(bold=True, size=14)
    lines = [
        "A row is written for every NSE session from 12 Oct 2026.",
        "Gap prediction is locked at 08:00 IST from the completed S&P session only.",
        "UP or DOWN is used only when that US move is at least 1%. Otherwise the row is NO CALL.",
        "NIFTY open and close are filled after they exist. They never change the prediction.",
        "Gap outcome compares the open with the previous close. That is the test.",
        "Open-to-close outcome is recorded separately. On past data that result was about a coin toss.",
        "Same way / Opposite are counted only on UP and DOWN rows. NO CALL rows are not scored.",
        "What the result shows is written after the open or the close. It explains the outcome. It does not change the locked call.",
        "This is a paper record. It is not a broker order and not a promise of profit.",
    ]
    for i, line in enumerate(lines, 3):
        about.cell(i, 1, line)
    about.column_dimensions["A"].width = 110
    tmp = path.with_suffix(".tmp.xlsx")
    book.save(tmp)
    tmp.replace(path)


def _write_site(payload: Dict[str, Any]) -> None:
    """Static page GitHub can host. The data is embedded, so the page needs no Python server."""
    root = Path(__file__).resolve().parents[3] / "docs"
    root.mkdir(parents=True, exist_ok=True)
    rows = payload.get("rows") or []
    summary = payload.get("summary") or {}
    latest = rows[-1] if rows else {}
    call = latest.get("gap_prediction") or "Waiting"
    body_rows = []
    for row in rows:
        body_rows.append(
            "<tr>"
            f"<td>{html.escape(str(row.get('session_date', '')))}</td>"
            f"<td>{html.escape(str(row.get('spx_move_pct', '')))}</td>"
            f"<td>{html.escape(str(row.get('gap_prediction', '')))}</td>"
            f"<td>{html.escape(str(row.get('gap_outcome', '')))}</td>"
            f"<td>{html.escape(str(row.get('open_to_close_outcome', '')))}</td>"
            f"<td>{html.escape(str(row.get('review', '')))}</td>"
            "</tr>"
        )
    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>NIFTY opening-gap call</title>
  <style>
    body {{ margin: 0; font-family: Georgia, "Times New Roman", serif; background: #0c1220; color: #e5e7eb; }}
    main {{ max-width: 980px; margin: 0 auto; padding: 32px 20px 64px; }}
    h1 {{ font-size: 32px; margin: 8px 0; }}
    p {{ line-height: 1.5; color: #cbd5e1; }}
    .call {{ font-size: 40px; margin: 12px 0; }}
    table {{ width: 100%; border-collapse: collapse; font-family: Arial, sans-serif; font-size: 14px; }}
    th, td {{ border-top: 1px solid #243044; text-align: left; padding: 10px 8px; vertical-align: top; }}
    th {{ color: #94a3b8; font-weight: 600; }}
    .muted {{ color: #94a3b8; font-size: 13px; }}
  </style>
</head>
<body>
<main>
  <p class="muted">Live test from 12 Oct 2026 · updated {html.escape(str(payload.get('updated_at', '')))}</p>
  <h1>NIFTY opening-gap call</h1>
  <p>One call is locked before 9:15 IST, and only when the completed S&amp;P move is at least 1%. The page is checked every 15 minutes until 16:00 IST so the open, the close, and a written reason can be saved. Those later checks do not make a new call. A one-minute or five-minute call is not used, because the measured link is the opening gap, not the rest of the session.</p>
  <p class="muted">{html.escape(str(latest.get('session_date', 'Next session')))} · {html.escape(str(latest.get('status', '')))}</p>
  <p class="call">{html.escape(str(call))}</p>
  <p>{html.escape(str(latest.get('review') or latest.get('notes') or ''))}</p>
  <p class="muted">Calls made {summary.get('calls', 0)} · No-call days {summary.get('no_calls', 0)} · Gap same way {summary.get('gap_same_way', 0)} of {summary.get('gap_scored', 0)}</p>
  <table>
    <thead><tr><th>Session</th><th>S&amp;P %</th><th>Call</th><th>Gap</th><th>Open to close</th><th>What the result shows</th></tr></thead>
    <tbody>
      {''.join(body_rows) or '<tr><td colspan="6">No sessions yet.</td></tr>'}
    </tbody>
  </table>
  <p class="muted">Paper record only. Past research, before 12 Oct 2026, is not this live test. Not investment advice.</p>
</main>
</body>
</html>
"""
    target = root / "index.html"
    tmp = root / "index.tmp.html"
    tmp.write_text(page, encoding="utf-8")
    tmp.replace(target)
    (root / "log.json").write_text(json.dumps(payload, default=str), encoding="utf-8")


def update_live_log(
    now: Optional[datetime] = None,
    path: Optional[Path] = None,
    nifty_bars: Optional[List[Dict[str, Any]]] = None,
    spx_bars: Optional[List[Dict[str, Any]]] = None,
    spx_stale: Optional[bool] = None,
    nifty_fetched_at: Optional[float] = None,
) -> Dict[str, Any]:
    """Refresh the Excel log and return the rows plus a short scoreboard."""
    clock = now or now_ist()
    if clock.tzinfo is None:
        raise ValueError("Live-log clock must include a timezone.")
    target = path or LOG_PATH.resolve()
    if nifty_bars is None or spx_bars is None or spx_stale is None:
        from app.data.market_history import get_daily_history, get_spx_history
        nifty = get_daily_history()
        spx = get_spx_history()
        nifty_bars = nifty.get("bars", [])
        spx_bars = spx.get("bars", [])
        spx_stale = bool(spx.get("is_stale", True))
        nifty_fetched_at = nifty.get("fetched_at")
    try:
        last = nse_session_to_describe(clock)
    except CalendarUnavailable as exc:
        return {"live_from": LIVE_START.isoformat(), "file": str(target), "rows": [], "summary": _summary([]), "warning": str(exc)}

    with _LOCK:
        existing = _read_rows(target)
        rows = []
        ordered = sorted(nifty_bars, key=lambda b: b.get("date", ""))
        for session in _sessions_through(last):
            row = existing.get(session.isoformat(), _blank_row(session))
            prev = [b for b in ordered if b.get("date", "") < session.isoformat()]
            _lock_prediction(row, session, spx_bars, bool(spx_stale), prev[-1]["close"] if prev else None)
            _score_row(row, session, ordered, clock, nifty_fetched_at)
            rows.append(row)
        unchanged = json.dumps(rows, sort_keys=True, default=str) == json.dumps(
            [existing[r["session_date"]] for r in rows if r["session_date"] in existing],
            sort_keys=True, default=str,
        ) and len(existing) == len(rows)
        if unchanged and target.exists():
            return {
                "live_from": LIVE_START.isoformat(),
                "file": str(target),
                "updated_at": clock.strftime("%Y-%m-%d %H:%M IST"),
                "rows": rows,
                "summary": _summary(rows),
                "unchanged": True,
            }
        try:
            _write(target, rows)
            warning = None
        except OSError as exc:
            warning = f"The Excel file could not be saved ({exc}). Close it if it is open, then refresh."
        payload = {
            "live_from": LIVE_START.isoformat(),
            "file": str(target),
            "updated_at": clock.strftime("%Y-%m-%d %H:%M IST"),
            "rows": rows,
            "summary": _summary(rows),
        }
        if warning:
            payload["warning"] = warning
        if path is None:
            _write_site(dict(payload, file="live_log/nifty_gap_log.xlsx"))
        return payload
